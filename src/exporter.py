"""
Експорт звітів у стилізований багатоаркушевий Excel-файл.
"""
import logging
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from src import config

logger = logging.getLogger(__name__)


def _style_header(ws: Worksheet, header_font: Font, header_fill: PatternFill) -> None:
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill


def _autofit_columns(ws: Worksheet) -> None:
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)


def export_report(
    master: pd.DataFrame,
    monthly: pd.DataFrame,
    product: pd.DataFrame,
    output_path: Path = config.OUTPUT_FILE,
) -> Path:
    """Записує три звіти у три аркуші одного Excel-файлу з корпоративним стилем."""
    header_font = Font(name=config.HEADER_FONT_NAME, size=11, bold=True,
                        color=config.HEADER_FONT_COLOR)
    header_fill = PatternFill(start_color=config.HEADER_FILL_COLOR,
                               end_color=config.HEADER_FILL_COLOR, fill_type="solid")
    regular_font = Font(name=config.HEADER_FONT_NAME, size=11)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        master.to_excel(writer, sheet_name="Cleaned Master Data", index=False)
        monthly.to_excel(writer, sheet_name="Monthly Performance", index=False)
        product.to_excel(writer, sheet_name="Product Performance", index=False)

        workbook = writer.book

        ws_master = workbook["Cleaned Master Data"]
        ws_master.freeze_panes = "A2"
        _style_header(ws_master, header_font, header_fill)
        for row in range(2, ws_master.max_row + 1):
            for col in range(1, 4):
                ws_master.cell(row=row, column=col).font = regular_font
            ws_master.cell(row=row, column=4).number_format = config.CURRENCY_FORMAT
            ws_master.cell(row=row, column=4).font = regular_font
            ws_master.cell(row=row, column=5).number_format = config.INTEGER_FORMAT
            ws_master.cell(row=row, column=5).font = regular_font
            ws_master.cell(row=row, column=6).number_format = config.CURRENCY_FORMAT
            ws_master.cell(row=row, column=6).font = regular_font

        ws_monthly = workbook["Monthly Performance"]
        _style_header(ws_monthly, header_font, header_fill)
        for row in range(2, ws_monthly.max_row + 1):
            ws_monthly.cell(row=row, column=1).font = regular_font
            ws_monthly.cell(row=row, column=2).number_format = config.INTEGER_FORMAT
            ws_monthly.cell(row=row, column=2).font = regular_font
            ws_monthly.cell(row=row, column=3).number_format = config.CURRENCY_FORMAT
            ws_monthly.cell(row=row, column=3).font = regular_font
            ws_monthly.cell(row=row, column=4).number_format = config.CURRENCY_FORMAT
            ws_monthly.cell(row=row, column=4).font = regular_font

        ws_product = workbook["Product Performance"]
        _style_header(ws_product, header_font, header_fill)
        for row in range(2, ws_product.max_row + 1):
            ws_product.cell(row=row, column=1).font = regular_font
            ws_product.cell(row=row, column=2).number_format = config.INTEGER_FORMAT
            ws_product.cell(row=row, column=2).font = regular_font
            ws_product.cell(row=row, column=3).number_format = config.CURRENCY_FORMAT
            ws_product.cell(row=row, column=3).font = regular_font

        for ws in (ws_master, ws_monthly, ws_product):
            _autofit_columns(ws)

    logger.info("Файл '%s' успішно збережено.", output_path)
    return output_path
