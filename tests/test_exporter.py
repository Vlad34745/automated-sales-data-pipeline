from pathlib import Path

import openpyxl
import pandas as pd

from src.exporter import export_report
from src import config


def make_sample_reports():
    master = pd.DataFrame({
        "Sale ID": [1, 2],
        "Date": ["2026-01-01", "2026-01-02"],
        "Product Name": ["A", "B"],
        "Unit Price": [10.0, 20.0],
        "Quantity": [1, 2],
        "Total Revenue": [10.0, 40.0],
    })
    monthly = pd.DataFrame({
        "year_month": ["2026-01"],
        "total_sales_volume": [3],
        "total_revenue_usd": [50.0],
        "average_ticket": [25.0],
    })
    product = pd.DataFrame({
        "product_name_cleaned": ["A", "B"],
        "units_sold": [1, 2],
        "revenue_generated_usd": [10.0, 40.0],
    })
    return master, monthly, product


def test_export_report_creates_file_at_given_path(tmp_path: Path):
    master, monthly, product = make_sample_reports()
    output_path = tmp_path / "report.xlsx"

    result_path = export_report(master, monthly, product, output_path)

    assert result_path == output_path
    assert output_path.exists()


def test_export_report_writes_all_three_sheets(tmp_path: Path):
    master, monthly, product = make_sample_reports()
    output_path = tmp_path / "report.xlsx"

    export_report(master, monthly, product, output_path)

    workbook = openpyxl.load_workbook(output_path)
    assert workbook.sheetnames == [
        "Cleaned Master Data", "Monthly Performance", "Product Performance",
    ]


def test_export_report_preserves_row_data(tmp_path: Path):
    master, monthly, product = make_sample_reports()
    output_path = tmp_path / "report.xlsx"

    export_report(master, monthly, product, output_path)

    written_master = pd.read_excel(output_path, sheet_name="Cleaned Master Data")
    assert written_master["Sale ID"].tolist() == [1, 2]
    assert written_master["Total Revenue"].tolist() == [10.0, 40.0]


def test_export_report_applies_header_style(tmp_path: Path):
    master, monthly, product = make_sample_reports()
    output_path = tmp_path / "report.xlsx"

    export_report(master, monthly, product, output_path)

    workbook = openpyxl.load_workbook(output_path)
    header_cell = workbook["Monthly Performance"]["A1"]
    assert header_cell.font.bold is True
    assert header_cell.fill.start_color.rgb.endswith(config.HEADER_FILL_COLOR)


def test_export_report_freezes_header_on_master_sheet(tmp_path: Path):
    master, monthly, product = make_sample_reports()
    output_path = tmp_path / "report.xlsx"

    export_report(master, monthly, product, output_path)

    workbook = openpyxl.load_workbook(output_path)
    assert workbook["Cleaned Master Data"].freeze_panes == "A2"


def test_export_report_applies_currency_format_to_revenue_column(tmp_path: Path):
    master, monthly, product = make_sample_reports()
    output_path = tmp_path / "report.xlsx"

    export_report(master, monthly, product, output_path)

    workbook = openpyxl.load_workbook(output_path)
    # 'Total Revenue' - шоста колонка на аркуші 'Cleaned Master Data'
    revenue_cell = workbook["Cleaned Master Data"].cell(row=2, column=6)
    assert revenue_cell.number_format == config.CURRENCY_FORMAT
